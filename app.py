"""
Skillhall - Campus skill exchange platform
Backend: Flask  |  Database: Excel workbook (data/skillhall.xlsx)

Run:   python app.py      then open http://127.0.0.1:5000
"""
import os
import threading
from datetime import datetime

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
DB_FILE = os.path.join(DATA_DIR, "skillhall.xlsx")
LOCK = threading.Lock()  # prevents two requests writing the file at once

# One sheet = one table. First row of each sheet = column names.
SHEETS = {
    "Students": ["id", "name", "department", "year", "email", "phone", "bio", "created_at"],
    "Skills": ["id", "student_id", "skill", "type", "level", "description", "created_at"],
    "Requests": ["id", "from_id", "to_id", "skill", "message", "status", "created_at"],
}

app = Flask(__name__)


# --------------------------------------------------------------------------
# Excel helpers
# --------------------------------------------------------------------------
def now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def init_db():
    os.makedirs(DATA_DIR, exist_ok=True)
    if os.path.exists(DB_FILE):
        return
    wb = Workbook()
    wb.remove(wb.active)
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1C2545")
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        for i, col in enumerate(cols):
            ws.column_dimensions[chr(65 + i)].width = max(14, len(col) + 6)
    wb.save(DB_FILE)


def read_rows(sheet):
    """Return every row of a sheet as a list of dicts."""
    wb = load_workbook(DB_FILE, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    header = rows[0]
    return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


def add_row(sheet, data):
    """Append a row; auto-generates the id. Returns the saved dict."""
    with LOCK:
        wb = load_workbook(DB_FILE)
        ws = wb[sheet]
        last_id = max([r[0] or 0 for r in ws.iter_rows(min_row=2, values_only=True)] or [0])
        data = {**data, "id": int(last_id) + 1}
        if "created_at" in SHEETS[sheet]:
            data["created_at"] = now()
        ws.append([data.get(c, "") for c in SHEETS[sheet]])
        wb.save(DB_FILE)
    return data


def update_row(sheet, row_id, changes):
    with LOCK:
        wb = load_workbook(DB_FILE)
        ws = wb[sheet]
        cols = SHEETS[sheet]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == row_id:
                for key, val in changes.items():
                    row[cols.index(key)].value = val
                wb.save(DB_FILE)
                return True
    return False


def delete_row(sheet, row_id):
    with LOCK:
        wb = load_workbook(DB_FILE)
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == row_id:
                ws.delete_rows(row[0].row)
                wb.save(DB_FILE)
                return True
    return False


def student_public(s):
    return {k: s[k] for k in ("id", "name", "department", "year", "email", "phone", "bio")}


@app.errorhandler(PermissionError)
def locked(_):
    return jsonify(error="The Excel file is open in another program. Close it and try again."), 423


# --------------------------------------------------------------------------
# Pages
# --------------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html")


# --------------------------------------------------------------------------
# Students (register / login / profile)
# --------------------------------------------------------------------------
@app.post("/api/register")
def register():
    d = request.get_json(force=True)
    name = (d.get("name") or "").strip()
    email = (d.get("email") or "").strip().lower()
    if not name or not email or "@" not in email:
        return jsonify(error="Name and a valid email are required."), 400
    if any((s["email"] or "").lower() == email for s in read_rows("Students")):
        return jsonify(error="This email is already registered. Use login instead."), 409
    s = add_row("Students", {
        "name": name,
        "department": (d.get("department") or "").strip(),
        "year": d.get("year") or "1st Year",
        "email": email,
        "phone": (d.get("phone") or "").strip(),
        "bio": (d.get("bio") or "").strip(),
    })
    return jsonify(student_public(s)), 201


@app.post("/api/login")
def login():
    email = (request.get_json(force=True).get("email") or "").strip().lower()
    for s in read_rows("Students"):
        if (s["email"] or "").lower() == email:
            return jsonify(student_public(s))
    return jsonify(error="No profile found for that email. Please register."), 404


@app.get("/api/students/<int:sid>")
def get_student(sid):
    for s in read_rows("Students"):
        if s["id"] == sid:
            skills = [k for k in read_rows("Skills") if k["student_id"] == sid]
            return jsonify({**student_public(s), "skills": skills})
    return jsonify(error="Student not found"), 404


@app.put("/api/students/<int:sid>")
def edit_student(sid):
    d = request.get_json(force=True)
    allowed = {k: d[k] for k in ("name", "department", "year", "phone", "bio") if k in d}
    if not update_row("Students", sid, allowed):
        return jsonify(error="Student not found"), 404
    return jsonify(ok=True)


# --------------------------------------------------------------------------
# Skills
# --------------------------------------------------------------------------
@app.post("/api/skills")
def add_skill():
    d = request.get_json(force=True)
    skill = (d.get("skill") or "").strip()
    if not skill or not d.get("student_id"):
        return jsonify(error="Skill name is required."), 400
    row = add_row("Skills", {
        "student_id": int(d["student_id"]),
        "skill": skill,
        "type": d.get("type") if d.get("type") in ("teach", "learn") else "teach",
        "level": d.get("level") or "Intermediate",
        "description": (d.get("description") or "").strip(),
    })
    return jsonify(row), 201


@app.delete("/api/skills/<int:skill_id>")
def remove_skill(skill_id):
    return jsonify(ok=delete_row("Skills", skill_id))


@app.get("/api/skills")
def search_skills():
    """Find students who TEACH a skill. ?q=python&me=3 (me is hidden from results)"""
    q = (request.args.get("q") or "").lower().strip()
    me = request.args.get("me", type=int)
    students = {s["id"]: s for s in read_rows("Students")}
    out = []
    for k in read_rows("Skills"):
        if k["type"] != "teach" or k["student_id"] == me or k["student_id"] not in students:
            continue
        s = students[k["student_id"]]
        hay = f'{k["skill"]} {k["description"]} {s["name"]} {s["department"]}'.lower()
        if q and q not in hay:
            continue
        out.append({**k, "student": student_public(s)})
    return jsonify(out)


# --------------------------------------------------------------------------
# Recommended matches
# --------------------------------------------------------------------------
@app.get("/api/matches/<int:sid>")
def matches(sid):
    skills = read_rows("Skills")
    students = {s["id"]: s for s in read_rows("Students")}
    my_teach = {k["skill"].lower() for k in skills if k["student_id"] == sid and k["type"] == "teach"}
    my_learn = {k["skill"].lower() for k in skills if k["student_id"] == sid and k["type"] == "learn"}
    result = []
    for oid, s in students.items():
        if oid == sid:
            continue
        theirs_teach = {k["skill"].lower() for k in skills if k["student_id"] == oid and k["type"] == "teach"}
        theirs_learn = {k["skill"].lower() for k in skills if k["student_id"] == oid and k["type"] == "learn"}
        they_teach_me = sorted(theirs_teach & my_learn)
        i_teach_them = sorted(theirs_learn & my_teach)
        if not they_teach_me:
            continue
        result.append({
            "student": student_public(s),
            "they_teach_me": they_teach_me,
            "i_teach_them": i_teach_them,
            "mutual": bool(i_teach_them),
            "score": len(they_teach_me) * 2 + len(i_teach_them) * 3,
        })
    result.sort(key=lambda r: r["score"], reverse=True)
    return jsonify(result)


# --------------------------------------------------------------------------
# Skill requests
# --------------------------------------------------------------------------
@app.post("/api/requests")
def create_request():
    d = request.get_json(force=True)
    if not d.get("from_id") or not d.get("to_id"):
        return jsonify(error="Missing student ids."), 400
    row = add_row("Requests", {
        "from_id": int(d["from_id"]),
        "to_id": int(d["to_id"]),
        "skill": d.get("skill", ""),
        "message": (d.get("message") or "").strip(),
        "status": "Pending",
    })
    return jsonify(row), 201


@app.get("/api/requests/<int:sid>")
def my_requests(sid):
    students = {s["id"]: s for s in read_rows("Students")}
    inbox, sent = [], []
    for r in read_rows("Requests"):
        if r["to_id"] == sid and r["from_id"] in students:
            inbox.append({**r, "other": student_public(students[r["from_id"]])})
        elif r["from_id"] == sid and r["to_id"] in students:
            sent.append({**r, "other": student_public(students[r["to_id"]])})
    return jsonify(inbox=inbox[::-1], sent=sent[::-1])


@app.put("/api/requests/<int:rid>")
def set_request_status(rid):
    status = request.get_json(force=True).get("status")
    if status not in ("Accepted", "Declined", "Pending"):
        return jsonify(error="Invalid status"), 400
    return jsonify(ok=update_row("Requests", rid, {"status": status}))


# --------------------------------------------------------------------------
# Dashboard + Database viewer + Excel download
# --------------------------------------------------------------------------
@app.get("/api/dashboard")
def dashboard():
    students, skills, reqs = read_rows("Students"), read_rows("Skills"), read_rows("Requests")
    counts = {}
    for k in skills:
        if k["type"] == "teach":
            key = k["skill"].strip().title()
            counts[key] = counts.get(key, 0) + 1
    top = sorted(counts.items(), key=lambda x: x[1], reverse=True)[:6]
    return jsonify({
        "students": len(students),
        "teach": sum(1 for k in skills if k["type"] == "teach"),
        "learn": sum(1 for k in skills if k["type"] == "learn"),
        "requests": len(reqs),
        "accepted": sum(1 for r in reqs if r["status"] == "Accepted"),
        "top_skills": [{"skill": s, "count": c} for s, c in top],
    })


@app.get("/api/database")
def database_view():
    return jsonify({name: {"columns": cols, "rows": read_rows(name)} for name, cols in SHEETS.items()})


@app.get("/api/download")
def download():
    return send_file(DB_FILE, as_attachment=True, download_name="skillhall.xlsx")


# --------------------------------------------------------------------------
# Demo data
# --------------------------------------------------------------------------
@app.post("/api/demo")
def demo():
    for s in read_rows("Students"):
        if s["email"] == "demo@skillhall.edu":
            return jsonify(student_public(s))
    people = [
        ("Demo Student", "Computer Science", "2nd Year", "demo@skillhall.edu", "9000000000",
         [("Excel", "teach", "Advanced"), ("Python", "learn", "Beginner"), ("Photoshop", "learn", "Beginner")]),
        ("Priya R", "Information Technology", "3rd Year", "priya@skillhall.edu", "9000000001",
         [("Python", "teach", "Advanced"), ("Excel", "learn", "Beginner")]),
        ("Arjun K", "Mechanical Engineering", "2nd Year", "arjun@skillhall.edu", "9000000002",
         [("Photoshop", "teach", "Intermediate"), ("Public Speaking", "learn", "Beginner")]),
        ("Meena S", "Electronics", "4th Year", "meena@skillhall.edu", "9000000003",
         [("Public Speaking", "teach", "Advanced"), ("Excel", "learn", "Intermediate")]),
    ]
    first = None
    for name, dept, year, email, phone, sk in people:
        s = add_row("Students", {"name": name, "department": dept, "year": year, "email": email,
                                 "phone": phone, "bio": "Happy to swap skills!"})
        first = first or s
        for skill, typ, lvl in sk:
            add_row("Skills", {"student_id": s["id"], "skill": skill, "type": typ, "level": lvl,
                               "description": f"{lvl} level {skill}"})
    return jsonify(student_public(first))


if __name__ == "__main__":
    init_db()
    app.run(debug=True)
